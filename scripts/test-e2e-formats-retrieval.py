from __future__ import annotations

import importlib.util
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from PIL import Image, ImageDraw, ImageFont
from docx import Document as DocxDocument
from pptx import Presentation

from rag3.haystack_native import ingest_source
from rag3.retrieval import search

COLLECTION = "rag3_haystack_e2e_formats"


def make_fixtures() -> list[Path]:
    d = ROOT / "artifacts" / "e2e-formats"
    d.mkdir(parents=True, exist_ok=True)

    txt = d / "e2e.txt"
    txt.write_text("RAG3 TXT evidence phrase: cobalt river alpha.\n", encoding="utf-8")

    docx = d / "e2e.docx"
    doc = DocxDocument()
    doc.add_heading("RAG3 DOCX Fixture", level=1)
    doc.add_paragraph("RAG3 DOCX evidence phrase: amber forest beta.")
    doc.save(docx)

    pptx = d / "e2e.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(1000000, 1000000, 7000000, 2000000)
    box.text_frame.text = "RAG3 PPTX evidence phrase: silver mountain gamma."
    prs.save(pptx)

    pdf = d / "e2e.pdf"
    pdf_text = "RAG3 PDF evidence phrase: violet ocean delta."
    objects = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    objects.append(
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
    )
    stream = f"BT /F1 18 Tf 72 720 Td ({pdf_text}) Tj ET".encode("ascii")
    objects.append(
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream + b"\nendstream"
    )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out.extend(f"{i} 0 obj\n".encode("ascii"))
        out.extend(obj)
        out.extend(b"\nendobj\n")
    xref = len(out)
    out.extend(f"xref\n0 {len(objects)+1}\n".encode("ascii"))
    out.extend(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        out.extend(f"{off:010d} 00000 n \n".encode("ascii"))
    out.extend(
        f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    pdf.write_bytes(bytes(out))

    img = d / "e2e.png"
    im = Image.new("RGB", (1800, 500), "white")
    draw = ImageDraw.Draw(im)
    font_path = Path(r"C:\Windows\Fonts\arial.ttf")
    font = ImageFont.truetype(str(font_path), 64) if font_path.exists() else ImageFont.load_default()
    draw.text((50, 180), "RAG3 IMAGE OCR evidence phrase: golden cloud epsilon.", fill="black", font=font)
    im.save(img)

    return [txt, docx, pptx, pdf, img]


def load_mcp_for_collection(collection: str):
    os.environ["RAG3_COLLECTION"] = collection
    path = ROOT / "scripts" / "rag3-mcp-server.py"
    spec = importlib.util.spec_from_file_location("rag3_mcp_server_e2e", path)
    if spec is None or spec.loader is None:
        raise RuntimeError("Could not load MCP module")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main() -> None:
    files = make_fixtures()
    results = []
    for i, p in enumerate(files):
        out = ingest_source(
            p,
            collection=COLLECTION,
            recreate_index=(i == 0),
        )
        results.append({
            "file": p.name,
            "documents_written": out["documents_written"],
            "count_after": out["count_after"],
            "source_preserved": Path(out["managed_source"]).is_file(),
        })

    direct = search(
        "Which fixture mentions violet ocean delta?",
        mode="hybrid",
        top_k=3,
        collection=COLLECTION,
    )
    if not direct["hits"]:
        raise SystemExit("Direct hybrid retrieval returned no hits")

    mcp_module = load_mcp_for_collection(COLLECTION)
    mcp_result = mcp_module.search_rag3(
        "Which fixture mentions amber forest beta?",
        mode="hybrid",
        top_k=3,
    )

    summary = {
        "formats": results,
        "direct_hit_count": len(direct["hits"]),
        "direct_top_title": direct["hits"][0]["title"],
        "mcp_hit_count": len(mcp_result["hits"]),
        "mcp_top_title": mcp_result["hits"][0]["title"] if mcp_result["hits"] else None,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))

    if not all(r["source_preserved"] and (r["documents_written"] or 0) >= 1 for r in results):
        raise SystemExit("One or more formats failed ingestion/source preservation")
    if len(mcp_result["hits"]) < 1:
        raise SystemExit("MCP retrieval returned no hits")


if __name__ == "__main__":
    main()
